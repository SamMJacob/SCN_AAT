"""Build the AAT report as a Word document: results/AAT_Report.docx.

Tables are read from the generated markdown files in results/ (so the numbers always match
`python -m src.eval.report`); the prose is the hand-checked discussion from REPORT_NOTES.md.
Run after the report step:  python scripts/make_report_docx.py
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results"
FIG = RES / "figures"
OUT = RES / "AAT_Report.docx"
ARCH_OUT = RES / "Architecture.docx"
DESKTOP = Path("/mnt/c/Users/sammj/OneDrive/Desktop")  # Windows desktop copy (WSL; Desktop is OneDrive-redirected)

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x52, 0x51, 0x4E)
HEAD_FILL, BAND_FILL, HYB_FILL = "1F3A5F", "F3F4F6", "E8F0FB"


# ---------------------------------------------------------------- markdown table helpers
def md_tables(path: Path) -> list[list[list[str]]]:
    """All pipe tables in a markdown file, each as rows of cells (header first)."""
    tables, cur = [], []
    for line in path.read_text().splitlines():
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                cur.append(cells)
        elif cur:
            tables.append(cur)
            cur = []
    if cur:
        tables.append(cur)
    return tables


def strip_md(s: str) -> str:
    return s.replace("**", "").replace("`", "")


def pick(rows, cols):
    idx = [rows[0].index(c) for c in cols]
    return [[r[i] for i in idx] for r in rows]


# ---------------------------------------------------------------- docx helpers
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self._styles()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.2))
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
        self._footer(sec)

    # styles
    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, before in (("Heading 1", 16, 18), ("Heading 2", 13, 12), ("Heading 3", 11, 10)):
            h = st[name]
            h.font.name = "Calibri"
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = NAVY
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            st[name].font.size = Pt(10.5)
            st[name].paragraph_format.space_after = Pt(3)
        cap = st["Caption"]
        cap.font.size = Pt(9)
        cap.font.italic = False
        cap.font.color.rgb = GREY
        cap.paragraph_format.space_after = Pt(10)

    def _footer(self, sec):
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("EV CAN-bus attack detection with ANFIS  ·  page ")
        r.font.size, r.font.color.rgb = Pt(8.5), GREY
        add_field(p, "PAGE")
        for r in p.runs:
            r.font.size, r.font.color.rgb = Pt(8.5), GREY

    # text
    def h(self, text, level=1):
        self.doc.add_heading(text, level)

    def p(self, text, style=None, italic=False, align=None):
        """Paragraph with **bold** and *italic* inline markup."""
        para = self.doc.add_paragraph(style=style)
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("`"):
                run = para.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
            elif part.startswith("**"):
                run = para.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("*"):
                run = para.add_run(part[1:-1])
                run.italic = True
            else:
                run = para.add_run(part)
            if italic:
                run.italic = True
        if align:
            para.alignment = align
        return para

    def bullets(self, items, numbered=False):
        for it in items:
            self.p(it, style="List Number" if numbered else "List Bullet")

    def callout(self, title, text):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        shade(c, HYB_FILL)
        cell_margins(t, 120, 120, 160, 160)
        para = c.paragraphs[0]
        r = para.add_run(title + "  ")
        r.bold = True
        r.font.color.rgb = NAVY
        for part in re.split(r"(\*\*[^*]+\*\*)", text):
            run = para.add_run(part.strip("*") if part.startswith("**") else part)
            run.bold = part.startswith("**")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # floats
    def figure(self, path, caption, width_cm=16.5):
        if not path.exists():
            self.p(f"[missing figure: {path.name}]", italic=True)
            return
        self.doc.add_picture(str(path), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        self.fig_n += 1
        cap = self.doc.add_paragraph(style="Caption")
        r = cap.add_run(f"Figure {self.fig_n}. ")
        r.bold = True
        cap.add_run(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, rows, caption, col_widths=None, font=8.5, highlight=None, bold_first_col=False):
        """rows: header + body (strings). highlight(row) -> True shades the row light blue."""
        if caption:
            self.tab_n += 1
            cap = self.doc.add_paragraph(style="Caption")
            r = cap.add_run(f"Table {self.tab_n}. ")
            r.bold = True
            cap.add_run(caption)
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_after = Pt(4)

        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        cell_margins(t)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                para = c.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                run = para.add_run(strip_md(val))
                run.font.size = Pt(font)
                if i == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    shade(c, HEAD_FILL)
                else:
                    if highlight and highlight(row):
                        shade(c, HYB_FILL)
                    elif i % 2 == 0:
                        shade(c, BAND_FILL)
                    if (bold_first_col and j == 0) or "**" in val:
                        run.bold = True
                if j > 0 and i > 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        repeat_header(t.rows[0])
        # light borders
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "D0D3D8")
            borders.append(el)
        tblPr.append(borders)
        if col_widths:
            for row in t.rows:
                for c, w in zip(row.cells, col_widths):
                    c.width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        self.doc.save(path)


# ---------------------------------------------------------------- content
def code_block(R, lines):
    t = R.doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, BAND_FILL)
    cell_margins(t, 80, 80, 120, 120)
    para = c.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    for i, ln in enumerate(lines):
        run = para.add_run(ln)
        run.font.name, run.font.size = "Consolas", Pt(7.5)
        if i < len(lines) - 1:
            run.add_break()
    R.doc.add_paragraph().paragraph_format.space_after = Pt(2)


def appendix_from_md(R, path: Path, letter: str | None = None, skip_title: bool = False):
    """Render a markdown file (headings, paragraphs, lists, tables, images, code) as an appendix."""
    lines = path.read_text().splitlines()
    i, para = 0, []

    def flush_para():
        if para:
            R.p(" ".join(para))
            para.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            flush_para()
            j = i + 1
            while not lines[j].strip().startswith("```"):
                j += 1
            code_block(R, lines[i + 1:j])
            i = j + 1
            continue
        if s.startswith("#"):
            flush_para()
            level = len(s) - len(s.lstrip("#"))
            title = s.lstrip("#").strip()
            if level == 1:
                if not skip_title:
                    R.h(f"Appendix {letter}. {title}" if letter else title, 1)
            elif letter:
                R.h(f"{letter}.{title}", min(level, 3))
            else:
                R.h(title, min(level - 1, 3))
            i += 1
            continue
        m = re.match(r"!\[(.*)\]\((.*)\)", s)
        if m:
            flush_para()
            R.figure((path.parent / m.group(2)).resolve(), m.group(1))
            i += 1
            continue
        if s.startswith("|"):
            flush_para()
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                j += 1
            rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines[i:j]]
            rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
            R.table(rows, None, font=8)
            i = j
            continue
        lm = re.match(r"(- |\d+\. )(.*)", s)
        if lm:
            flush_para()
            item = [lm.group(2)]
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() \
                    and not re.match(r"\s*(- |\d+\. )", lines[i]):
                item.append(lines[i].strip())
                i += 1
            text = " ".join(item)
            if lm.group(1) == "- ":
                R.p(text, style="List Bullet")
            else:      # manual numbers: Word's shared numbering would continue across lists
                q = R.p(lm.group(1) + text)
                q.paragraph_format.left_indent = Cm(0.63)
                q.paragraph_format.first_line_indent = Cm(-0.63)
                q.paragraph_format.space_after = Pt(3)
            continue
        if not s:
            flush_para()
        else:
            para.append(s)
        i += 1
    flush_para()


def is_hybrid(row):
    return row[0].startswith("ANFIS")


def build():
    R = Report()
    d = R.doc

    # ---------- title page
    for _ in range(5):
        d.add_paragraph()
    t = R.p("Detecting Cyber-Attacks on the EV CAN Bus with ANFIS", align=WD_ALIGN_PARAGRAPH.CENTER)
    t.runs[0].font.size, t.runs[0].bold, t.runs[0].font.color.rgb = Pt(24), True, NAVY
    s = R.p("A reproduction and extension: hybrid vs non-hybrid soft computing, "
            "and the effect of clustering-based initialisation", align=WD_ALIGN_PARAGRAPH.CENTER)
    s.runs[0].font.size, s.runs[0].font.color.rgb = Pt(13), GREY
    d.add_paragraph()
    R.p("Soft Computing — Alternative Assessment Task (Hybrid Systems)", align=WD_ALIGN_PARAGRAPH.CENTER)
    a = R.p("Sam M Jacob", align=WD_ALIGN_PARAGRAPH.CENTER)
    a.runs[0].bold = True
    R.p(date.today().strftime("%d %B %Y"), align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(6):
        d.add_paragraph()
    R.p("Reproduced paper: O. A. Al Isawi, K. A. Al Jaafari, A. S. Al-Sumaiti, “Electric Vehicles CAN Bus "
        "Cyber Attacks Detection Using Adaptive Neuro Fuzzy Inference System,” *IEEE Access*, vol. 13, "
        "pp. 90862–90874, 2025. DOI 10.1109/ACCESS.2025.3550970.", align=WD_ALIGN_PARAGRAPH.CENTER)
    d.paragraphs[-1].runs[0].font.size = Pt(9)
    R.page_break()

    # ---------- abstract
    R.h("Abstract")
    R.p("Al Isawi et al. report that an Adaptive Neuro-Fuzzy Inference System (ANFIS), initialised by "
        "subtractive clustering (SC) or fuzzy c-means (FCM) and stacked with a second ANFIS, detects DoS and "
        "Impersonation attacks on the OTIDS CAN-bus dataset with ≈99.7 % accuracy. We reimplemented the full "
        "pipeline in Python (ANFIS from scratch in PyTorch, hybrid least-squares + gradient learning) and "
        "evaluated it under one protocol shared by every model: train-only scaling and SMOTE, validation-only "
        "tuning and five random seeds.")
    R.p("The paper's numbers **reproduce** (our ANFIS-FCM+stk: 99.99 % binary, 100 % multiclass accuracy), "
        "but the data analysis shows why: the class label is a function of the CAN ID, which is also an "
        "input. On a fairer benchmark that removes the ID (**A-noID**), a **Random Forest is the best "
        "model** (macro-F1 0.990 / 0.989). The hybrid ANFIS clearly beats its own non-hybrid ingredients: "
        "+0.036 / +0.101 macro-F1 over a pure fuzzy (FCM) classifier, and +0.019 over an MLP in the "
        "binary task.")
    R.p("For initialisation, a clustering start beats random placement, and SC is the most reliable choice "
        "because it sizes the rule base itself. However, the ranking between clustering methods depends on "
        "the rule count and largely disappears after stacking. Finally, re-labelling Impersonation by attack "
        "time shows that the attacker's frames are statistically indistinguishable from the genuine ECU's. "
        "The paper's Impersonation result therefore measures ‘is this ID 0x164?’ rather than attack "
        "detection.")

    # ---------- 1 introduction
    R.h("1. Introduction")
    R.p("The Controller Area Network (CAN) carries every control message inside a vehicle but has no "
        "authentication: any node can transmit any identifier. Intrusion-detection systems (IDS) therefore "
        "try to recognise malicious frames from their identifier, payload and timing. Al Isawi et al. apply "
        "ANFIS, a hybrid that combines a Takagi–Sugeno fuzzy rule base (interpretable IF–THEN rules) with "
        "neural-network-style learning of its parameters, to the OTIDS dataset. They compare it with "
        "classical classifiers and deep-learning baselines.")
    R.p("This project reproduces that study and uses it to answer two questions about hybrid soft computing:")
    R.bullets([
        "**Q1. Non-hybrid vs hybrid.** Does the neuro-fuzzy hybrid outperform a pure neural model (MLP), a "
        "pure fuzzy model (FCM classifier) and classical models (LDA, LR, linear/RBF SVM, Random Forest)?",
        "**Q2. Initialisation.** How much does the clustering method used to build the initial fuzzy "
        "inference system matter: Random vs K-Means (hard) vs Subtractive Clustering vs Fuzzy C-Means (fuzzy)?",
    ])
    R.p("The rules were fixed before any results: nothing is tuned on the test set, the hybrid used for "
        "the headline comparison (ANFIS-FCM+stk, the paper's best architecture) was chosen in advance, and "
        "negative results are reported as found.")

    # ---------- 2 data
    R.h("2. Dataset and what it actually contains")
    R.p("OTIDS (HCRL, Korea University) contains four CAN logs recorded from a KIA Soul: attack-free "
        "traffic, a DoS attack (flooding with the highest-priority ID 0x000), a fuzzy attack and an "
        "impersonation attack on ID 0x164. Each line holds a timestamp, CAN ID, DLC and up to eight data "
        "bytes. All 4.6 M lines were parsed without loss.")
    ov = md_tables(RES / "data_report.md")[0]
    ov = [["log", "rows", "duration (s)", "frames/s", "unique IDs", "0x000 frames", "0x164 frames"]] + \
        [[r[0], f"{int(r[1]):,}", f"{float(r[2]):.0f}", f"{float(r[3]):.0f}", r[4], f"{int(r[7]):,}",
          f"{int(r[8]):,}"] for r in ov[1:]]
    R.table(ov, "The four OTIDS logs.", font=9)

    R.h("2.1 Labelling findings", 2)
    R.p("The paper labels all 0x000 frames as DoS and **every** 0x164 frame as Impersonation. Before "
        "designing experiments we checked whether these labels correspond to something visible in the frames:")
    R.bullets([
        "**DoS is genuinely distinct.** The DoS log contains 335,624 injected 0x000 frames with a single "
        "all-zero payload, in bursts (median gap 0.5 ms), throughout the log. However, all-zero payloads also "
        "occur on normal IDs (159,104 such frames in the attack-free log), so without the ID, DoS must be "
        "recognised partly by timing.",
        "**Impersonation is invisible.** In the impersonation log, 0x164 traffic before and after the "
        "documented attack start (250 s) has the same rate (119.2 vs 118.9 Hz), the same payload patterns "
        "(χ² = 3.6, p = 0.999; no post-250 s payload is new), the same rolling-counter behaviour, the same "
        "total bus load, and no ID changes rate by more than 3 %. Only a tiny inter-arrival shift exists (KS "
        "D = 0.068 vs D = 0.016 for the same comparison on attack-free data).",
        "**Consequence.** The paper's Impersonation class is defined by the ID alone: genuine and "
        "malicious 0x164 frames share one label. With the CAN ID as an input feature, the task reduces to "
        "an ID lookup.",
    ])
    R.figure(FIG / "data_impersonation_0x164.png",
             "0x164 traffic in the impersonation log before vs after the 250 s attack start: message rate, "
             "inter-arrival time and payload-pattern frequency are indistinguishable.")
    R.figure(FIG / "data_dos_0x000.png", "DoS log: 0x000 injection runs over the whole recording.", 14)

    # ---------- 3 method
    R.h("3. Method")
    R.h("3.1 Experimental settings", 2)
    R.p("Because of the findings above, the study runs four settings. **A-noID is the main benchmark** for "
        "both research questions.")
    R.table([
        ["Setting", "Labels", "Inputs", "Purpose"],
        ["A", "by CAN ID (paper)", "CAN ID, DLC, d0–d7", "Reproduce the paper"],
        ["A-noID", "by CAN ID (paper)", "DLC, d0–d7, log10 Δt (same ID)",
         "Main benchmark: label-leaking ID removed"],
        ["B", "Imp = 0x164 after 250 s", "as A-noID", "Is the true attack learnable?"],
        ["B-ID", "as B", "as A-noID + CAN ID", "As B, with the ID"],
    ], "Experimental settings.", col_widths=[2, 4, 5, 5.5], font=9, bold_first_col=True)

    R.h("3.2 Protocol (identical for every model)", 2)
    R.bullets([
        "**Sample:** 20,000 frames per seed with the paper's class totals (1,497 DoS / 979 Imp / 17,524 "
        "Normal), stratified 80/20 train/test split.",
        "**Preprocessing:** min-max scaling fitted on the training set only, then SMOTE (k = 30) applied to "
        "the training set only.",
        "**Tuning:** baseline hyperparameters are selected on a validation split (20 % of training, carved out "
        "before SMOTE). ANFIS uses the paper's hyperparameters unchanged. The test set is never used for "
        "any choice.",
        "**Seeds:** 5 (each re-draws subsample, split and SMOTE). Reported as mean ± std. The primary metric "
        "is macro-F1, since accuracy is dominated by the 88 % Normal class.",
        "**Significance:** paired per-seed differences. A difference counts as ‘within noise’ if |mean Δ| ≤ "
        "its std over seeds or |Δ| < 0.0025 (≈ one Impersonation test error). An exact McNemar test is run "
        "on seed 0.",
    ])

    R.h("3.3 Models", 2)
    R.table([
        ["Family", "Model", "Notes"],
        ["Hybrid (neuro-fuzzy)", "ANFIS (OvR)",
         "First-order Takagi–Sugeno, Gaussian MFs. One ANFIS per class (one-vs-rest)."],
        ["", "ANFIS + stacking (+stk)",
         "Paper's design: a second ‘augmented’ ANFIS takes the stage-1 class scores and regresses the class index."],
        ["Non-hybrid soft computing", "MLP", "Pure neural network (sklearn, early stopping)."],
        ["", "FCM classifier", "Pure fuzzy: FCM clusters → membership-weighted class vote."],
        ["Classical ML", "LDA, LR, linear SVM, RBF SVM, Random Forest", "Validation-tuned."],
    ], "Model ladder.", col_widths=[4, 4.5, 8], font=9)
    R.p("**ANFIS implementation.** Rules have Gaussian membership functions per input and linear "
        "consequents. Firing strengths are computed in log space and normalised. Training follows Jang's "
        "hybrid rule: in each epoch, the consequent parameters are solved by (ridge) least squares and the "
        "premise centres and widths are updated by one Adam pass. The consequents are re-solved after the "
        "final epoch. The initial FIS comes from clustering the joint [input, target] space (as MATLAB's "
        "genfis2/genfis3 do). Initialisations:")
    R.bullets([
        "**Random:** R = 10 centres drawn at random from the training points; widths = global feature std.",
        "**K-Means:** R = 10 hard clusters; widths = per-cluster standard deviation.",
        "**Subtractive clustering (SC):** Chiu's method with the paper's radius and squash factor; the "
        "number of rules is decided by the algorithm.",
        "**FCM:** R = 10 fuzzy clusters (m = 2); widths = membership-weighted standard deviation.",
    ])
    R.p("Unit tests check the parser, train-only preprocessing, subtractive clustering on synthetic blobs, "
        "ANFIS fitting sin·cos and XOR, and agreement of the least-squares step with numpy (13 tests, all pass).")

    # ---------- 4 results
    R.page_break()
    R.h("4. Results")
    R.h("4.1 Reproducing the paper (Setting A)", 2)
    notes = md_tables(RES / "REPORT_NOTES.md")
    rep = notes[0]
    R.table(rep, "Setting A: paper vs our reproduction (accuracy and macro-F1, mean ± std over 5 seeds).",
            col_widths=[1.6, 4.4, 1.9, 3.0, 2.2, 3.0], font=8.5, highlight=lambda r: "ANFIS" in r[1])
    R.p("The ANFIS results are reproduced and slightly exceeded. However, 11 of 15 models reach macro-F1 "
        "≥ 0.99 in multiclass, including the MLP, RF, RBF-SVM and our validation-tuned LR. The paper's weak "
        "linear baselines are explained by the ID: 0x164 = 356 sits in the middle of the ID range, so one "
        "linear threshold cannot isolate it (LR on the ID alone: macro-F1 0.718). With the regularisation "
        "chosen on the validation split, LR reaches 0.996. The paper's LR (0.866) matches an untuned, "
        "strongly regularised fit. **The ≈99.7 % headline is a property of the setting, not evidence that "
        "the hybrid is needed.**")
    R.figure(FIG / "ladder_multi.png",
             "Model ladder (multiclass macro-F1) across settings. Blue = hybrid ANFIS, orange = non-hybrid "
             "soft computing, green = classical. In Setting A nearly everything sits at the ceiling.")

    # 4.2 comparison 1
    R.h("4.2 Comparison 1: non-hybrid vs hybrid (A-noID)", 2)
    for task, label in (("binary", "binary (Attack vs Normal)"), ("multi", "multiclass (DoS / Imp / Normal)")):
        rows = md_tables(RES / f"summary_A-noID_{task}.md")[0]
        cols = ["model", "family", "macro-F1", "accuracy"] + \
            (["F1 Attack"] if task == "binary" else ["F1 DoS", "F1 Imp"]) + ["rules", "train s"]
        rows = pick(rows, cols)
        R.table(rows, f"A-noID, {label}: test metrics, mean ± std over 5 seeds, sorted by macro-F1. "
                      "Hybrid rows shaded.", font=7.5, highlight=is_hybrid)
    R.figure(FIG / "comparison1_A-noID_multi.png",
             "Comparison 1, A-noID multiclass: macro-F1 by model family (mean ± std, 5 seeds).", 15)

    c1 = notes[1]
    c1 = [r for r in c1 if r[0] in ("setting", "A-noID", "A")]
    c1 = pick(c1, ["setting", "task", "best non-hybrid", "MLP", "FCM-clf", "ANFIS-FCM+stk",
                   "best ANFIS (post-hoc)", "McNemar p (seed 0)"])
    R.table(c1, "Headline comparison (macro-F1). ANFIS-FCM+stk was pre-specified; ‘best ANFIS’ is chosen on "
                "test and therefore optimistic. McNemar compares best non-hybrid vs ANFIS-FCM+stk.", font=8)
    R.table([
        ["A-noID", "vs Random Forest", "vs MLP (pure neural)", "vs FCM classifier (pure fuzzy)"],
        ["binary", "−0.010 ± 0.005 (0/5 seeds)\nhybrid worse, McNemar p ≈ 1e-7",
         "+0.019 ± 0.018 (4/5)\nhybrid better", "+0.036 ± 0.007 (5/5)\nhybrid better"],
        ["multi", "−0.004 ± 0.006 (1/5)\nwithin noise", "+0.002 ± 0.008 (4/5)\nwithin noise",
         "+0.101 ± 0.018 (5/5)\nhybrid better"],
    ], "Paired differences for ANFIS-FCM+stk (Δ macro-F1, positive = hybrid better; seeds won).",
        col_widths=[2, 5, 4.5, 5], font=8.5, bold_first_col=True)
    R.bullets([
        "**The best overall model is non-hybrid: Random Forest.** Even the best ANFIS picked post hoc "
        "(ANFIS-SC 0.984 binary; ANFIS-KMeans+stk 0.987 multi) does not beat it.",
        "**Against its own ingredients the hybrid wins clearly.** Adding neural learning (least-squares "
        "consequents + gradient-tuned membership functions) to FCM clusters is the largest single "
        "improvement in the study. The gain over the MLP is modest, and the MLP is unstable across seeds "
        "(std 0.017 in binary).",
        "**Where the errors are.** Almost all remaining errors are Normal frames predicted as DoS (seed 0: "
        "RF 19, ANFIS-FCM+stk 17 of 3,505). These are normal frames with an all-zero payload, which only "
        "timing can separate once the ID is removed. Impersonation is recognised perfectly by both models.",
        "**Cost.** ANFIS trains in 1–11 s with 20–63 rules, while RF needs under 1 s. ANFIS's advantage is "
        "a compact, inspectable rule base, not accuracy.",
    ])
    R.figure(FIG / "confusion_A-noID_multi.png",
             "Seed-0 confusion matrices, A-noID multiclass: best non-hybrid (RF) vs ANFIS-FCM+stk.", 13)
    R.callout("Answer to Q1.",
              "Hybridisation helps **relative to its components**: a large gain over the pure fuzzy model and "
              "a modest, task-dependent gain over a pure MLP. It does **not** beat a strong classical "
              "ensemble (Random Forest) on this data.")

    # 4.3 comparison 2
    R.h("4.3 Comparison 2: clustering initialisation (A-noID)", 2)
    R.p("Only the initial FIS changes. Data, scaling, SMOTE, training (1 epoch), OvR and stacking are fixed. "
        "K-Means vs FCM isolates the value of *fuzzy* (partial) memberships over hard assignment at the "
        "same rule count.")
    c2 = notes[2]
    c2 = [r for r in c2 if r[0] in ("setting", "A-noID", "A")]
    R.table(c2, "Macro-F1 (OvR / +stk), total rule count and training time per initialisation "
                "(R / K / SC / F = Random / K-Means / SC / FCM).", font=7.5)
    R.figure(FIG / "comparison2_A-noID_multi.png",
             "Comparison 2, A-noID multiclass: macro-F1, Impersonation F1, rule count and training time per "
             "initialisation, unstacked (blue) vs stacked (orange).")
    abl = md_tables(RES / "ablations.md")[0]
    r20 = {(r[1], r[2]): r[6] for r in abl[1:] if r[0] == "A-noID" and r[3] == "R20"}
    main = {(r[1], r[2]): r[5] for r in abl[1:] if r[0] == "A-noID" and r[3] == "R20"}
    sc = {t: md_tables(RES / f"comparison2_clustering_A-noID_{t}.md")[0] for t in ("binary", "multi")}
    sc_val = {t: next(r[1] for r in sc[t][1:] if r[0] == "ANFIS-SC") for t in sc}
    rows = [["init", "binary, R = 10", "binary, R = 20", "multi, R = 10", "multi, R = 20"]]
    for init in ("Random", "KMeans", "FCM"):
        m = f"ANFIS-{init}"
        rows.append([init, main[("binary", m)], r20[("binary", m)], main[("multi", m)], r20[("multi", m)]])
    rows.append(["SC (self-sized: 43–50 rules)", sc_val["binary"], "–", sc_val["multi"], "–"])
    R.table(rows, "Rule-count confound: unstacked macro-F1 with R = 10 vs R = 20 rules per OvR member.",
            font=8.5, bold_first_col=True)
    R.bullets([
        "**At the paper's R = 10**, SC is the best unstacked init (0.984 / 0.983), beating Random on 5/5 "
        "seeds in both tasks with the lowest variance. K-Means beats Random in multiclass (+0.014, 5/5). FCM "
        "is within noise of Random and *below* K-Means in multiclass (−0.007, 0/5 seeds).",
        "**But the ranking is confounded with capacity.** SC chooses more rules (43–50 vs 20–30). At R = 20, "
        "FCM rises to 0.971 / 0.982 and Random to 0.977 / 0.967, while K-Means does not improve. The FCM vs "
        "K-Means verdict flips with R. On the validation split, R = 5 is clearly worse and R = 10 ≈ R = 20.",
        "**Stacking evens out the inits:** K-Means/SC/FCM+stk all land at 0.985–0.987 (multi) and "
        "0.979–0.984 (binary). Only Random+stk lags (0.966 multi).",
        "**Epochs:** 5 epochs (early stopping on validation, then refit) change macro-F1 by less than the "
        "seed noise for every init, as the paper states. Early stopping picks 1 epoch most often for SC.",
        "**Time:** FCM initialisation makes ANFIS-FCM 4–6× slower than K-Means/SC (8.9 s vs 1.4 / 2.0 s, "
        "multiclass).",
    ])
    R.figure(FIG / "mf_overlay_A-noID_multi.png",
             "Membership functions of one OvR member on one input, before (dashed) and after (solid) hybrid "
             "training, per initialisation. This is the interpretable part of the hybrid.")
    R.callout("Answer to Q2.",
              "A clustering initialisation beats random placement at the paper's rule count, and **SC is the "
              "most reliable choice** because it also selects a sensible rule count. Differences between "
              "clustering methods are small, depend on R and mostly disappear once the stacking ANFIS is "
              "added. We find **no consistent evidence that fuzzy clustering (FCM) beats hard clustering "
              "(K-Means)** as an ANFIS initialiser.")

    # 4.4 setting B
    R.h("4.4 Setting B: is the real Impersonation attack learnable?", 2)
    R.p("Setting B labels only 0x164 frames sent after the attack start as Impersonation; genuine 0x164 "
        "frames before 250 s are Normal. Every model (classical, MLP, all ANFIS variants) reaches "
        "Impersonation F1 ≈ 0.79. This equals the trivial rule ‘Imp iff ID = 0x164’ (0.794 ± 0.018) on the "
        "same test sets. The models learn to recognise 0x164-*type* messages, not the attacker.")
    pb = md_tables(RES / "probe_B.md")[0]
    keep = [pb[0]] + [r for r in pb[1:] if r[1] != "all"]
    keep = [[r[0].replace(" (real boundary)", "").replace(" (fake boundary, control)", " (control)"),
             r[1], r[2], r[3]] for r in keep]
    R.table(keep, "Learnability probe: classify each 0x164 frame as after vs before 250 s. Time-blocked "
                  "5-fold CV (10 s groups), ROC AUC per fold. Control = same probe on attack-free data with "
                  "a fake boundary. AUC 0.5 = no signal.",
            font=8, highlight=lambda r: r[2] == "RF" and "timing" in r[1] and "control" not in r[0])
    R.p("Payload carries no signal (AUC 0.50 everywhere). Only a non-linear model on timing finds a faint "
        "fingerprint (RF AUC 0.59–0.61 vs 0.50 on the control). **Per-frame detection of OTIDS "
        "impersonation is essentially not possible from these features**, so the paper's Impersonation "
        "results measure ID recognition.")

    # 4.5 ablations
    R.h("4.5 Ablations", 2)
    sel = []
    for setting, task, model, var in [
        ("A-noID", "binary", "ANFIS-FCM+stk", "oof"), ("A-noID", "multi", "ANFIS-FCM+stk", "oof"),
        ("A-noID", "multi", "ANFIS-SC+stk", "oof"),
        ("A-noID", "binary", "FCM-clf", "nosmote"), ("A-noID", "multi", "FCM-clf", "nosmote"),
        ("A-noID", "multi", "RF", "nosmote"), ("A-noID", "multi", "ANFIS-FCM", "nosmote"),
        ("A-noID", "multi", "ANFIS-FCM+stk", "nosmote"),
        ("A-noID", "multi", "ANFIS-SC", "ep5"), ("A-noID", "multi", "ANFIS-FCM", "ep5"),
    ]:
        r = next(x for x in abl[1:] if x[:4] == [setting, task, model, var])
        sel.append([r[1], r[2], {"oof": "OOF stacking", "nosmote": "no SMOTE", "ep5": "5 epochs"}[r[3]],
                    r[5], r[6], r[7]])
    R.table([["task", "model", "variant", "main", "variant", "Δ (variant − main)"]] + sel,
            "Selected ablations on A-noID (macro-F1, paired over 5 seeds). Full list in results/ablations.md.",
            font=8)
    R.bullets([
        "**Paper (in-sample) vs out-of-fold stacking:** OOF is *worse* for K-Means/FCM and equal for SC. A "
        "10-rule stage-1 ANFIS barely overfits, so its training outputs resemble its test outputs. OOF "
        "stage 2 is instead trained on noisier scores than it sees at test time.",
        "**SMOTE** is essential for the pure fuzzy classifier (−0.14 / −0.30 without it), because the "
        "cluster class distributions get swamped by Normal. RF is unaffected, and unstacked ANFIS even "
        "improves without it.",
        "**Removing the CAN ID** (A → A-noID) costs RF only 0.010, ANFIS 0.012–0.050 and linear models up "
        "to 0.14.",
    ])

    # ---------- 5 conclusions
    R.h("5. Discussion and conclusions")
    R.bullets([
        "**The paper reproduces, but its headline is uninformative.** With the ID as input and labels "
        "defined by the ID, almost any non-linear model reaches ≈100 %.",
        "**Hybrid vs non-hybrid:** ANFIS is a genuine improvement over the pure fuzzy model it is built from "
        "and is competitive with (sometimes better than) an MLP, with a far smaller, interpretable rule base. "
        "It is not the most accurate model: a Random Forest is better or equal everywhere.",
        "**Initialisation:** any clustering start is better than random. SC is the safest default because "
        "it sets the number of rules from the data. Hard vs fuzzy clustering makes no consistent difference "
        "once rule count and stacking are accounted for.",
        "**Dataset validity:** OTIDS Impersonation, as labelled by the paper, is ID recognition. A real "
        "impersonation detector would need sequence/timing models and a dataset in which the attack "
        "leaves a trace.",
    ])

    R.h("6. Deviations from the paper and limitations")
    R.bullets([
        "The paper does not state which files its 20,000 rows came from, and its class ratios match no log. "
        "We drew DoS from the DoS log, Impersonation from the impersonation log and Normal from the "
        "attack-free log (other IDs).",
        "ANFIS was reimplemented in PyTorch rather than MATLAB. The stacking ANFIS regresses the class index "
        "(our reading of an ambiguous description).",
        "The paper tuned on its test set; we tune only on a validation split.",
        "The paper's deep-learning baselines (CNN, LSTM, GRU, CNN-LSTM) were not re-run; their reported "
        "numbers are quoted only.",
        "Five seeds give limited power. ‘Within noise’ is a seed-spread rule, not a formal test. McNemar "
        "is run on seed 0 only, without multiple-comparison correction.",
        "‘Best ANFIS (post-hoc)’ is selected on test and is optimistic; the headline uses the "
        "pre-specified ANFIS-FCM+stk.",
        "Training times are single-machine wall-clock (ANFIS on GPU, others on CPU) and indicative only.",
    ])

    R.h("References")
    R.bullets([
        "O. A. Al Isawi, K. A. Al Jaafari, A. S. Al-Sumaiti, “Electric Vehicles CAN Bus Cyber Attacks "
        "Detection Using Adaptive Neuro Fuzzy Inference System,” IEEE Access, 13, 90862–90874, 2025.",
        "H. Lee, S. H. Jeong, H. K. Kim, “OTIDS: A Novel Intrusion Detection System for In-vehicle Network by "
        "Using Remote Frame,” PST 2017. Dataset: HCRL CAN-intrusion dataset.",
        "J.-S. R. Jang, “ANFIS: Adaptive-Network-Based Fuzzy Inference System,” IEEE Trans. SMC, 23(3), "
        "665–685, 1993.",
        "S. L. Chiu, “Fuzzy Model Identification Based on Cluster Estimation,” J. Intelligent & Fuzzy "
        "Systems, 2(3), 267–278, 1994.",
        "J. C. Bezdek, Pattern Recognition with Fuzzy Objective Function Algorithms, Plenum, 1981.",
        "N. V. Chawla et al., “SMOTE: Synthetic Minority Over-sampling Technique,” JAIR, 16, 321–357, 2002.",
    ], numbered=True)
    R.p("Code, configuration and all result tables: project repository (README.md, results/).", italic=True)

    R.page_break()
    appendix_from_md(R, ROOT / "docs" / "ARCHITECTURE.md", "A")

    save_both(R, OUT)


def save_both(R, out):
    for path in (out, DESKTOP / out.name):
        if path.parent.is_dir():
            R.save(path)
            print(f"wrote {path}  ({R.fig_n} figures, {R.tab_n} tables)")
        else:
            print(f"skipped {path} (folder not found)")


def build_architecture():
    """Standalone architecture document from docs/ARCHITECTURE.md."""
    R = Report()
    d = R.doc
    for _ in range(6):
        d.add_paragraph()
    t = R.p("System Architecture and Implementation", align=WD_ALIGN_PARAGRAPH.CENTER)
    t.runs[0].font.size, t.runs[0].bold, t.runs[0].font.color.rgb = Pt(24), True, NAVY
    s = R.p("EV CAN-bus attack detection with ANFIS: data pipeline, hybrid neuro-fuzzy models, "
            "baselines and evaluation", align=WD_ALIGN_PARAGRAPH.CENTER)
    s.runs[0].font.size, s.runs[0].font.color.rgb = Pt(13), GREY
    d.add_paragraph()
    R.p("Soft Computing — Alternative Assessment Task (Hybrid Systems)", align=WD_ALIGN_PARAGRAPH.CENTER)
    a = R.p("Sam M Jacob", align=WD_ALIGN_PARAGRAPH.CENTER)
    a.runs[0].bold = True
    R.p(date.today().strftime("%d %B %Y"), align=WD_ALIGN_PARAGRAPH.CENTER)
    R.page_break()
    appendix_from_md(R, ROOT / "docs" / "ARCHITECTURE.md", skip_title=True)
    save_both(R, ARCH_OUT)


if __name__ == "__main__":
    build()
    build_architecture()
